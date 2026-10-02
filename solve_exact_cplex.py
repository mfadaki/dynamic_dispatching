"""
solve_exact_cplex.py
======================
Solves the EXACT dispatching MDP (full state enumeration, one LP variable
per state -- no VFA, no basis functions, no sampling) via CPLEX, and saves
every state's exact value V*(s) to exact_values.csv (and .xlsx), plus every
(state,action) pair's exact Q(s,a) to q_values.csv (and .xlsx, when it fits
Excel's row limit -- see save_q_values), for later structural-property
checks (queue-balancing dominance, submodularity, etc.) against ground
truth rather than an approximation.

Output folder: results/<YYYYmmdd-HHMMSS>_{exact}/ -- same naming as this
project's other runs (date-time, then the method in braces, e.g.
20260911-040524_{exact}), containing exact_values.csv/.xlsx,
q_values.csv/.xlsx, and inputs_exact.py: a byte-for-byte copy of the
config as it was when this run STARTED (captured at startup, not at the
end, so editing inputs_exact.py while a multi-hour solve is running cannot
change what gets recorded).

State enumeration, the exact one-step transition model, and the LP
constraint construction (rhs_const, trans) below are lifted VERBATIM from
evaluate_approximation.py's own enumerate_states()/exact_transitions()/
_state_key()/_snap_tau() -- not reimplemented -- specifically so this
script's state space and dynamics are identical to the ones
evaluate_approximation.py already validates elsewhere in this project. The
only thing that changes is which solver turns those constraints into V*:
CPLEX here, scipy.optimize.linprog(method='highs') there.

Files this script depends on (and nothing else from the project):
    inputs/inputs_exact.py     -- the instance parameters
    classes/mdp_exact.py       -- the MDP dynamics, importing its parameters
                                  directly from inputs/inputs_exact.py
NO other config module is used or imported anywhere in this exact-solve
path: there is no command-line config option, no alias, and no
sys.modules substitution. classes/mdp_exact.py is a copy of
classes/mdp_evaluate_approximation.py whose only functional difference is
that one import line.

classes/alp_evaluate_approximation.py is NOT needed -- this script never
touches phi()/theta at all, since it computes the exact value function
directly, with no approximation architecture in the loop. The basis-function
fields in inputs_exact.py (NO_BASIS_FN, THETA_*) are therefore never read.

Usage
-----
    python solve_exact_cplex.py

CPLEX licensing note
---------------------
The pip-installable `cplex` package defaults to the free Community
Edition, which is HARD-CAPPED at 1000 variables and 1000 constraints
(confirmed directly: solves fine at exactly 1000, raises "CPLEX Error
1016: Community Edition. Problem size limits exceeded" at 1001). The exact
LP here has one variable per state and (states x actions) constraints, so
any instance with more than ~333 states (at 3 actions) will exceed the
Community Edition limit. If you see Error 1016, you need a full
academic/commercial CPLEX license properly configured (e.g. via the IBM
Academic Initiative), not a code fix. This script detects that specific
error and falls back to scipy's HiGHS solver automatically, WITH A LOUD,
explicit warning and a record of which solver actually produced the saved
values -- it does not silently substitute a different solver.
"""

import sys
import os
import time
import itertools
from datetime import datetime
import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def _stage(label):
    print(f"\n=== {label}  [{time.strftime('%H:%M:%S')}] ===", flush=True)


_last_progress_print = {}


def _progress(done, total, t_start, label, min_interval_s=2.0):
    """Prints at most once every min_interval_s seconds per label (always
    prints on completion), so a five-second stage doesn't flood the
    terminal and a twenty-minute stage still gives regular, readable
    updates -- what's done so far, and an ETA for what's left."""
    now = time.time()
    is_done = (done >= total)
    if not is_done and (now - _last_progress_print.get(label, 0.0)) < min_interval_s:
        return
    _last_progress_print[label] = now
    elapsed = now - t_start
    frac = done / total if total else 1.0
    rate = done / elapsed if elapsed > 0 else 0.0
    eta = (total - done) / rate if rate > 0 else float('nan')
    print(f"  [{label}] {done:,}/{total:,} ({frac*100:5.1f}%)  "
          f"elapsed={elapsed:6.1f}s  ETA={eta:6.1f}s", flush=True)


_CONFIG_MODULE_NAME = 'inputs_exact'      # fixed: the only config this script uses

# Drop any cached copies first, so a Python session that already imported
# these modules (e.g. an interactive kernel where you edited inputs_exact.py
# and re-ran) picks up the CURRENT files rather than stale ones.
for _m in ('inputs.inputs_exact', 'classes.mdp_exact'):
    sys.modules.pop(_m, None)

import inputs.inputs_exact as cfg                       # noqa: E402
import classes.mdp_exact as _mdp_module                 # noqa: E402
from classes.mdp_exact import MDP                       # noqa: E402

_MDP_PARAM_NAMES = [
    'GAMMA', 'N_LABS', 'LAB_IDS', 'L_AGE', 'TAU_MAX', 'DELTA_T',
    'LAMBDA_AGE', 'LAMBDA', 'MU', 'LAMBDA_TOTAL', 'C_DISPATCH', 'H_HOLD',
    'C_EXP_DEPOT', 'C_EXP_LAB', 'K_CAPACITY', 'N_MAX', 'N_MIN', 'N_INV',
    'N_STATE', 'STATE_BOUNDS', 'ACTION_SET', 'ACTION_BOUNDS',
]


def _verify_mdp_uses_inputs_exact():
    """Cheap safeguard: every parameter the MDP module holds must equal
    inputs_exact's. True by construction now (classes/mdp_exact.py imports
    straight from inputs_exact), but it stops the run loudly if that module
    is ever edited to read parameters from somewhere else, instead of
    silently solving a different instance under inputs_exact's name."""
    bad = []
    for name in _MDP_PARAM_NAMES:
        mdp_val = np.asarray(getattr(_mdp_module, name), dtype=float)
        cfg_val = np.asarray(getattr(cfg, name), dtype=float)
        if mdp_val.shape != cfg_val.shape or not np.array_equal(mdp_val, cfg_val):
            bad.append(f"  {name}: MDP module has {getattr(_mdp_module, name)!r}, "
                       f"inputs_exact has {getattr(cfg, name)!r}")
    if bad:
        raise RuntimeError(
            "The MDP is NOT using inputs_exact.py's parameters -- refusing "
            "to solve a different instance than the one requested:\n"
            + "\n".join(bad)
        )


_verify_mdp_uses_inputs_exact()


def _capture_inputs_exact_source():
    """Read inputs_exact.py's bytes NOW (so the copy saved with the results
    is exactly what this run used, even if the file is edited during the
    solve) and check that this source text reproduces the parameters that
    were actually loaded. Python reuses cached bytecode when a file's mtime
    (1 s resolution) and size are unchanged, so a fast same-size edit can
    leave the run using OLD values while the file on disk shows NEW ones --
    this stops that at startup, before hours of solving, rather than saving
    a copy that disagrees with the results."""
    path = cfg.__file__
    with open(path, 'rb') as f:
        raw = f.read()
    ns = {'__name__': 'inputs_exact_source_check', '__file__': path}
    exec(compile(raw, path, 'exec'), ns)
    bad = []
    for name in _MDP_PARAM_NAMES + ['epochs_per_day']:
        src_val = np.asarray(ns[name], dtype=float)
        run_val = np.asarray(getattr(cfg, name), dtype=float)
        if src_val.shape != run_val.shape or not np.array_equal(src_val, run_val):
            bad.append(f"  {name}: file on disk gives {ns[name]!r}, loaded module has {getattr(cfg, name)!r}")
    if bad:
        raise RuntimeError(
            "inputs_exact.py on disk does not match the parameters actually "
            "loaded (stale cached bytecode). Delete the inputs/__pycache__ "
            "folder and re-run:\n" + "\n".join(bad)
        )
    return raw


_INPUTS_EXACT_SOURCE = _capture_inputs_exact_source()


def _make_results_dir(method='exact'):
    """results/<YYYYmmdd-HHMMSS>_{<method>}/ -- the same naming as this
    project's saveResultsFn (timestamp taken when the results are saved),
    e.g. 20260911-040524_{exact}. Anchored to this script's own folder, not
    the current working directory. Also writes the inputs_exact.py copy."""
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           'results', f"{stamp}_{{{method}}}")
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, 'inputs_exact.py'), 'wb') as f:
        f.write(_INPUTS_EXACT_SOURCE)
    print(f"  Results folder: {out_dir}", flush=True)
    return out_dir


# =============================================================================
# State enumeration and exact transitions -- VERBATIM from
# evaluate_approximation.py (see that file for the authoritative version;
# duplicated here, not imported, because importing evaluate_approximation.py
# directly would trigger ITS OWN config-loading side effects at import time).
# =============================================================================

def enumerate_states(mdp, cfg):
    n_cells = cfg.N_INV
    cell_values = range(0, cfg.N_MAX + 1)
    tau_grid = [cfg.TAU_MAX - k * cfg.DELTA_T for k in range(cfg.epochs_per_day)]

    total_combos = len(cell_values) ** n_cells
    t0 = time.time()
    states = []
    for ci, combo in enumerate(itertools.product(cell_values, repeat=n_cells)):
        n = np.array(combo, dtype=float).reshape(cfg.N_LABS + 1, cfg.L_AGE)
        for tau in tau_grid:
            states.append(mdp.build_state(n, tau))
        _progress(ci + 1, total_combos, t0, "enumeration")
    states = np.array(states)

    index = {_state_key(s, tau_grid): i for i, s in enumerate(states)}
    return states, index, tau_grid


def _snap_tau(tau, tau_grid, tol=1e-6):
    diffs = [abs(tau - g) for g in tau_grid]
    j = int(np.argmin(diffs))
    if diffs[j] > tol:
        raise KeyError(
            f"tau={tau} is not within {tol} of any enumerated grid value "
            f"{tau_grid} — check DELTA_T/TAU_MAX consistency."
        )
    return tau_grid[j]


def _state_key(s, tau_grid, ndigits=6):
    n_part = tuple(round(float(v), ndigits) for v in s[:-1])
    tau_part = round(_snap_tau(float(s[-1]), tau_grid), ndigits)
    return n_part + (tau_part,)


def exact_transitions(mdp, s, a, index, tau_grid):
    out = []
    for event in range(mdp.N_EVENT_TYPES):
        p = mdp._probs[event]
        if p <= 0:
            continue
        s_next = mdp._transition_single(s, a, event)
        key = _state_key(s_next, tau_grid)
        if key not in index:
            raise KeyError(
                "Transition left the enumerated state space -- K_CAPACITY "
                "too small for the arrival/processing rates chosen."
            )
        out.append((index[key], p))
    return out


# =============================================================================
# Exact LP via CPLEX (with an explicit, logged scipy fallback if the
# Community Edition size limit is hit)
# =============================================================================

def solve_exact_lp_cplex(mdp, states, index, tau_grid):
    """Builds the SAME exact LP as evaluate_approximation.py's
    solve_exact_lp (one variable per state, constraints
    V(s_i) - gamma*sum_j P(j|i,a)*V(s_j) <= r(s_i,a) for every (i,a)),
    solved via CPLEX instead of scipy. Returns (V_star, Q_values,
    solver_used) -- solver_used is 'cplex' or 'scipy_highs_fallback',
    recorded explicitly rather than left ambiguous.
    """
    M = len(states)
    actions = mdp.action_set
    n_actions = len(actions)

    _stage(f"LP construction ({M:,} states x {n_actions} actions "
          f"= {M * n_actions:,} constraints)")

    rhs_const = np.zeros((M, n_actions))
    trans = [[None] * n_actions for _ in range(M)]
    t0 = time.time()
    for i, s in enumerate(states):
        for ai, a in enumerate(actions):
            rhs_const[i, ai] = mdp._cost_single_raw(s, a) / mdp.cost_scale
            trans[i][ai] = exact_transitions(mdp, s, a, index, tau_grid)
        _progress(i + 1, M, t0, "LP construction")

    _stage("Solving")
    try:
        V_star = _solve_with_cplex(mdp, M, n_actions, rhs_const, trans)
        solver_used = 'cplex'
    except Exception as e:
        msg = str(e)
        if 'Error  1016' in msg or 'Error 1016' in msg or 'Community Edition' in msg:
            print("\n" + "!" * 70)
            print("CPLEX COMMUNITY EDITION SIZE LIMIT HIT (Error 1016).")
            print(f"This instance has {M} variables / {M*n_actions} constraints, "
                  f"exceeding the free Community Edition's 1000/1000 cap.")
            print("FALLING BACK TO scipy.optimize.linprog(method='highs') -- "
                  "this is NOT CPLEX. If you need a genuine CPLEX-solved "
                  "result at this instance size, a full academic/commercial "
                  "CPLEX license must be properly configured first.")
            print("!" * 70 + "\n")
            V_star = _solve_with_scipy(M, rhs_const, trans, mdp.gamma)
            solver_used = 'scipy_highs_fallback'
        else:
            raise

    Q_values = _compute_q_values(mdp, M, n_actions, rhs_const, trans, V_star)
    return V_star, Q_values, solver_used


def _compute_q_values(mdp, M, n_actions, rhs_const, trans, V_star):
    """Q(s,a) = r(s,a) + gamma*E[V*(s')|s,a], computed directly from the
    SAME rhs_const/trans already built for the LP above -- no re-querying
    the MDP. In the SAME normalized units as V_star (both come from the
    same normalized LP), so min_a Q(s,a) == V*(s) exactly at every state,
    by construction of the exact-LP characterization of V* (the LP's own
    constraints are V(s) <= r(s,a)+gamma*E[V(s')|s,a] for every a, tight
    at the optimal action) -- checked below directly, not just asserted."""
    _stage(f"Computing Q-values ({M:,} states x {n_actions} actions)")
    Q = np.zeros((M, n_actions))
    t0 = time.time()
    for i in range(M):
        for ai in range(n_actions):
            cont = sum(p * V_star[j] for j, p in trans[i][ai])
            Q[i, ai] = rhs_const[i, ai] + mdp.gamma * cont
        _progress(i + 1, M, t0, "Q-value computation")

    max_gap = float(np.max(np.abs(Q.min(axis=1) - V_star)))
    print(f"  Verification: max|min_a Q(s,a) - V*(s)| = {max_gap:.2e} "
          f"across all {M:,} states -- "
          f"{'PASS' if max_gap < 1e-6 else 'FAIL, something is wrong -- do not trust these Q-values'}",
          flush=True)
    return Q


def _solve_with_cplex(mdp, M, n_actions, rhs_const, trans):
    import cplex

    c = cplex.Cplex()
    # Native CPLEX iteration log left ON (to stdout) rather than
    # suppressed -- this is the only way to see CPLEX's own solve
    # actively progressing (simplex/barrier iteration count, objective
    # value) once c.solve() is called below; nothing outside CPLEX's own
    # C-level solve loop can print into it. Warning/results streams stay
    # suppressed -- they're post-solve diagnostics, not progress signal.
    c.set_log_stream(sys.stdout)
    c.set_warning_stream(None)
    c.set_results_stream(None)

    # V(s_i), i=0..M-1, unbounded (both directions) -- V* can be negative
    # (these are costs, not rewards, under this project's sign convention).
    c.variables.add(obj=[1.0 / M] * M, lb=[-1e20] * M, ub=[1e20] * M)
    # Maximize mean(V) subject to V(s) <= r(s,u)+gamma*E[V(s')|s,u] for
    # every (s,u) -- the standard exact-LP characterization of V*.
    c.objective.set_sense(c.objective.sense.maximize)

    _stage("Building CPLEX constraint rows")
    lin_expr, senses, rhs = [], [], []
    t0 = time.time()
    for i in range(M):
        for ai in range(n_actions):
            # Accumulate coefficients per column index before handing rows
            # to CPLEX -- unlike scipy's sparse matrix (A_ub[row,j] += ...,
            # which auto-combines), CPLEX rejects a row with the same
            # column index listed twice. This matters concretely here: a
            # "no-op" event can transition state i back to itself (j==i),
            # which would otherwise collide with the +1.0 V(s_i) term and
            # raise "CPLEX Error 1222: Duplicate entry or entries."
            # (confirmed directly -- this is not a hypothetical edge case).
            coeffs = {i: 1.0}
            for j, p in trans[i][ai]:
                coeffs[j] = coeffs.get(j, 0.0) - mdp.gamma * p
            idx = list(coeffs.keys())
            val = list(coeffs.values())
            lin_expr.append([idx, val])
            senses.append('L')
            rhs.append(float(rhs_const[i, ai]))
        _progress(i + 1, M, t0, "CPLEX row-building")

    c.linear_constraints.add(lin_expr=lin_expr, senses=senses, rhs=rhs)

    _stage("CPLEX solve starting -- CPLEX's own iteration log follows below")
    t_solve = time.time()
    c.solve()
    print(f"  CPLEX solve finished in {time.time()-t_solve:.1f}s", flush=True)

    status = c.solution.get_status_string()
    if 'optimal' not in status.lower():
        raise RuntimeError(f"CPLEX did not reach an optimal solution: {status}")

    return np.array(c.solution.get_values())


def _solve_with_scipy(M, rhs_const, trans, gamma):
    from scipy.optimize import linprog
    from scipy.sparse import lil_matrix

    n_actions = rhs_const.shape[1]
    n_constraints = M * n_actions
    A_ub = lil_matrix((n_constraints, M))
    b_ub = np.zeros(n_constraints)

    _stage("Building scipy sparse constraint matrix")
    t0 = time.time()
    row = 0
    for i in range(M):
        for ai in range(n_actions):
            A_ub[row, i] += 1.0
            for j, p in trans[i][ai]:
                A_ub[row, j] += -gamma * p
            b_ub[row] = rhs_const[i, ai]
            row += 1
        _progress(i + 1, M, t0, "scipy matrix-building")

    A_ub = A_ub.tocsr()
    c_obj = -np.ones(M) / M

    _stage("scipy HiGHS solve starting -- HiGHS's own iteration log follows below")
    t_solve = time.time()
    res = linprog(c_obj, A_ub=A_ub, b_ub=b_ub,
                  bounds=[(None, None)] * M, method='highs',
                  options={'disp': True})
    print(f"  scipy HiGHS solve finished in {time.time()-t_solve:.1f}s", flush=True)
    if not res.success:
        raise RuntimeError(f"scipy fallback LP failed to solve: {res.message}")
    return res.x


def save_q_values(states, Q_values, cfg, solver_used, out_dir):
    """Writes q_values.csv (and .xlsx): one row per (state, action) pair
    -- LONG format, M*n_actions rows, not M rows like exact_values.csv/
    .xlsx. Same state columns as exact_values.csv, plus 'action' (raw
    int, matching cfg.ACTION_SET) and 'action_label' (human-readable:
    'hold' / '->lab{p}', matching the action_label() convention already
    used by this project's other diagnostic scripts, e.g.
    diagnose_policy_disagreement.py)."""
    os.makedirs(out_dir, exist_ok=True)
    M, n_actions = Q_values.shape
    n_rows = M * n_actions

    col_names = []
    for p in range(cfg.N_LABS + 1):
        loc = 'depot' if p == 0 else f'lab{p}'
        for a in range(cfg.L_AGE):
            col_names.append(f'{loc}_age{a+1}')
    col_names += ['tau', 'action', 'Q_value']

    _stage(f"Building Q-value table ({M:,} states x {n_actions} actions "
          f"= {n_rows:,} rows) -- vectorized, not a Python loop")
    t0 = time.time()
    n_inv = cfg.N_INV
    state_part = np.concatenate([states[:, :n_inv], states[:, n_inv:n_inv + 1]], axis=1)
    state_rows = np.repeat(state_part, n_actions, axis=0)       # each state's row repeated n_actions times, in order
    action_col = np.tile(np.arange(n_actions), M).reshape(-1, 1)
    q_col = Q_values.reshape(-1, 1)                             # row-major flatten matches np.repeat's row order above
    rows = np.concatenate([state_rows, action_col, q_col], axis=1)
    action_labels = ['hold' if a == 0 else f'->lab{a}' for a in np.tile(np.arange(n_actions), M)]
    print(f"  Built {n_rows:,}-row table in {time.time()-t0:.1f}s", flush=True)

    csv_path = os.path.join(out_dir, 'q_values.csv')
    _stage(f"Writing Q-value CSV ({n_rows:,} rows)")
    t0 = time.time()
    header = ','.join(col_names)
    np.savetxt(csv_path, rows, delimiter=',', header=header, comments='',
              fmt='%.10g')
    print(f"  Wrote {csv_path} in {time.time()-t0:.1f}s "
          f"({n_rows:,} rows, solver={solver_used}). Numeric-only "
          f"('action' is the raw int, e.g. 0=hold/1=->lab1/2=->lab2 -- see "
          f"cfg.ACTION_SET); the human-readable action_label is in the "
          f".xlsx version only, where the write is already row-by-row so "
          f"adding a text column costs nothing extra there.", flush=True)

    EXCEL_ROW_LIMIT = 1_048_576
    if n_rows + 1 > EXCEL_ROW_LIMIT:   # +1 for the header row
        print(f"  SKIPPING .xlsx for Q-values: {n_rows:,} rows + header "
              f"exceeds Excel's hard limit of {EXCEL_ROW_LIMIT:,} rows per "
              f"sheet. q_values.csv above is complete and has everything; "
              f"use that instead.", flush=True)
        return

    try:
        import openpyxl
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Q Values"
        ws.append(col_names + ['action_label'])
        _stage(f"Writing Q-value Excel ({n_rows:,} rows -- row-by-row, "
              f"this is the slowest export step: {n_actions}x the row "
              f"count of exact_values.xlsx)")
        t0 = time.time()
        for i, (r, lbl) in enumerate(zip(rows, action_labels)):
            ws.append(list(r) + [lbl])
            _progress(i + 1, n_rows, t0, "Q-value Excel row-writing")
        xlsx_path = os.path.join(out_dir, 'q_values.xlsx')
        wb.save(xlsx_path)
        print(f"  Wrote {xlsx_path}", flush=True)
    except ImportError:
        print("openpyxl not available -- q_values.xlsx not written, "
              "q_values.csv is still complete.")


# =============================================================================
# CSV / Excel export
# =============================================================================

def save_state_values(states, V_star, cfg, solver_used, out_dir):
    """Writes exact_values.csv (and .xlsx) with one row per state: every
    inventory cell broken out by (location, age), tau, and V*(s). Location
    0 = depot, 1..N_LABS = labs, matching this project's n_{p,a} indexing
    throughout."""
    os.makedirs(out_dir, exist_ok=True)

    col_names = []
    for p in range(cfg.N_LABS + 1):
        loc = 'depot' if p == 0 else f'lab{p}'
        for a in range(cfg.L_AGE):
            col_names.append(f'{loc}_age{a+1}')
    col_names += ['tau', 'V_star']

    n_inv = cfg.N_INV
    rows = np.concatenate([states[:, :n_inv], states[:, n_inv:n_inv+1],
                           V_star.reshape(-1, 1)], axis=1)

    csv_path = os.path.join(out_dir, 'exact_values.csv')
    header = ','.join(col_names)
    _stage(f"Writing CSV ({len(states):,} rows)")
    t0 = time.time()
    np.savetxt(csv_path, rows, delimiter=',', header=header, comments='',
              fmt='%.10g')
    print(f"  Wrote {csv_path} in {time.time()-t0:.1f}s "
          f"({len(states)} states, solver={solver_used})", flush=True)

    try:
        import openpyxl
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Exact Values"
        ws.append(col_names)
        _stage(f"Writing Excel ({len(states):,} rows -- row-by-row, "
              f"typically much slower than the CSV write above)")
        t0 = time.time()
        for i, row in enumerate(rows):
            ws.append(list(row))
            _progress(i + 1, len(rows), t0, "Excel row-writing")
        ws2 = wb.create_sheet("Run Info")
        ws2.append(["config_module", _CONFIG_MODULE_NAME])
        ws2.append(["n_states", len(states)])
        ws2.append(["solver_used", solver_used])
        ws2.append(["N_LABS", cfg.N_LABS])
        ws2.append(["L_AGE", cfg.L_AGE])
        ws2.append(["K_CAPACITY", cfg.K_CAPACITY])
        ws2.append(["MU", str(list(cfg.MU))])
        ws2.append(["H_HOLD_lab1", str(list(cfg.H_HOLD[1]))])
        ws2.append(["H_HOLD_lab2", str(list(cfg.H_HOLD[2])) if cfg.N_LABS >= 2 else "n/a"])
        xlsx_path = os.path.join(out_dir, 'exact_values.xlsx')
        wb.save(xlsx_path)
        print(f"Wrote {xlsx_path}")
    except ImportError:
        print("openpyxl not available -- .xlsx not written, .csv is still complete.")


# =============================================================================
# Entry point
# =============================================================================

if __name__ == "__main__":
    _t_script_start = time.time()
    print(f"Config: inputs/{_CONFIG_MODULE_NAME}.py (the only config this script uses)")
    print(f"N_LABS={cfg.N_LABS}, L_AGE={cfg.L_AGE}, K_CAPACITY={cfg.K_CAPACITY}, "
          f"epochs_per_day={cfg.epochs_per_day}")
    print(f"MU={list(cfg.MU)}  (symmetric labs check: "
          f"{'PASS' if cfg.N_LABS < 2 or len(set(np.round(cfg.MU,10))) == 1 else 'FAIL -- MU differs across labs'})")

    mdp = MDP()
    mdp.calibrate_cost_scale(n_samples=getattr(cfg, 'NORM_N_SAMPLES', 5000),
                             seed=getattr(cfg, 'NORM_SEED', 12345))

    _stage("Enumerating states")
    states, index, tau_grid = enumerate_states(mdp, cfg)
    print(f"  Enumerated {len(states):,} states "
          f"(total elapsed so far: {time.time()-_t_script_start:.1f}s)", flush=True)

    V_star, Q_values, solver_used = solve_exact_lp_cplex(mdp, states, index, tau_grid)
    print(f"Solved via: {solver_used}")
    print(f"E[V*] (uniform over states) = {V_star.mean():.4f}")

    _stage("Saving results")
    out_dir = _make_results_dir('exact')
    save_state_values(states, V_star, cfg, solver_used, out_dir)
    save_q_values(states, Q_values, cfg, solver_used, out_dir)

    _stage(f"Done -- total wall-clock time: {time.time()-_t_script_start:.1f}s")